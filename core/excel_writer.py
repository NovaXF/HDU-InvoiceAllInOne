from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from models.invoice import Stack


def generate_stack_excel(
    stack: Stack,
    date_str: str,
    output_dir: Path,
    reimburser: str = "",
    stack_dir: Path | None = None,
) -> Path:
    """为一叠发票生成报销 Excel 表格，返回文件路径。"""
    wb = Workbook()
    ws = wb.active

    folder_name = stack.folder_name(date_str)
    ws.title = "报销清单"

    headers = ["序号", "项目名称", "发票代码", "发票号码", "发票金额", "实际报销人"]
    col_widths = [6, 40, 24, 24, 12, 14]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    title_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    title_font = Font(bold=True, size=12)
    header_font = Font(bold=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=folder_name)
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = center

    for i, invoice in enumerate(stack.invoices, 1):
        row = i + 2
        item_text = "\n".join(invoice.item_names) if invoice.item_names else ""

        values = [
            i,
            item_text,
            invoice.invoice_code or "",
            invoice.invoice_number or "",
            invoice.amount if invoice.amount is not None else "",
            invoice.remark,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            cell.alignment = left_wrap if col == 2 else center

    empty_rows = max(0, 6 - len(stack.invoices))
    for j in range(empty_rows):
        row = len(stack.invoices) + 3 + j
        ws.cell(row=row, column=1, value=len(stack.invoices) + 1 + j)
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.alignment = center

    stack_dir = stack_dir or output_dir / folder_name
    stack_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = stack_dir / f"{folder_name}.xlsx"
    wb.save(xlsx_path)
    return xlsx_path
